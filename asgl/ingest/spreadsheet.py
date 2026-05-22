from dataclasses import dataclass
from typing import Any, List, Tuple, Dict
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from asgl.utils.logging import setup_logging

log = setup_logging("asgl.ingest")

@dataclass
class TableExtract:
    sheet: str
    bbox: Tuple[int,int,int,int]
    dataframe: pd.DataFrame
    score: float
    signals: Dict[str, Any]

def _is_blank(v):
    return v is None or (isinstance(v,str) and v.strip()=="")

def _parse_number(x):
    if x is None: return None
    if isinstance(x,(int,float,np.number)): return float(x)
    if isinstance(x,str):
        s=x.strip()
        if s in {"-","—","–"}: return None
        s=s.replace("£","").replace("$","").replace(",","")
        neg=False
        if s.startswith("(") and s.endswith(")"):
            neg=True; s=s[1:-1]
        if s.endswith("%"):
            try: return float(s[:-1])/100.0
            except: return None
        try:
            v=float(s)
            return -v if neg else v
        except: 
            return None
    return None

def _score_table(df:pd.DataFrame)->float:
    if df.empty: return 0.0
    density = df.notna().sum().sum() / (df.shape[0]*df.shape[1] + 1e-9)
    # header quality: how many text in first row
    first=df.iloc[0].astype(str)
    text_ratio = sum([1 for v in first if v and v!="None" and not v.replace(".","",1).isdigit()]) / max(1,len(first))
    return float(density*0.7 + text_ratio*0.3)

def extract_tables_xlsx(path: str, max_tables:int=5)->List[TableExtract]:
    wb = load_workbook(path, data_only=True)
    extracts=[]
    for ws in wb.worksheets:
        # used range
        cells=[]
        for row in ws.iter_rows():
            for c in row:
                if not _is_blank(c.value) and not (c.data_type=="n" and c.value==0):
                    cells.append((c.row,c.column))
        if not cells:
            continue
        rows=[r for r,_ in cells]; cols=[c for _,c in cells]
        min_r,max_r=min(rows),max(rows); min_c,max_c=min(cols),max(cols)
        grid=[]
        for r in range(min_r,max_r+1):
            grid.append([ws.cell(r,c).value for c in range(min_c,max_c+1)])
        df=pd.DataFrame(grid)
        # naive segmentation: split on long blank rows
        blank_rows=[i for i in range(len(df)) if df.iloc[i].isna().all()]
        splits=[-1]+blank_rows+[len(df)]
        for a,b in zip(splits[:-1], splits[1:]):
            if b-a<=2: 
                continue
            block=df.iloc[a+1:b].copy()
            if block.notna().sum().sum()<10:
                continue
            # header index = first row in block
            headers=block.iloc[0].astype(str).tolist()
            body=block.iloc[1:].copy()
            body.columns=[h if h and h!="None" else f"col_{j}" for j,h in enumerate(headers)]
            for col in body.columns:
                body[col]=body[col].apply(_parse_number)
            sc=_score_table(block)
            signals={"used_range":[min_r,min_c,max_r,max_c],"block_rows":int(body.shape[0]),"block_cols":int(body.shape[1]),"density":float(block.notna().sum().sum()/(block.size))}
            extracts.append(TableExtract(ws.title,(min_r,min_c,max_r,max_c),body,sc,signals))
    extracts=sorted(extracts, key=lambda t: t.score, reverse=True)[:max_tables]
    log.info("xlsx_tables_extracted", extra={"extra":{"path":path,"tables":len(extracts)}})
    return extracts
