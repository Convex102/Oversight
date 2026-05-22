
import sys, json, uuid
from PyQt6.QtWidgets import QApplication, QWidget, QVBoxLayout, QPushButton, QFileDialog, QLabel, QLineEdit, QHBoxLayout, QListWidget
from PyQt6.QtCore import Qt
from openpyxl import load_workbook

class App(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("ASGL Template Picker (Local)")
        self.wb=None; self.path=None; self.sheet=None
        self.mappings=[]
        layout=QVBoxLayout()
        self.info=QLabel("Load a working papers .xlsx then map key cells.")
        layout.addWidget(self.info)
        btn=QPushButton("Load .xlsx"); btn.clicked.connect(self.load); layout.addWidget(btn)
        self.template=QLineEdit(); self.template.setPlaceholderText("Template name"); layout.addWidget(self.template)
        row=QHBoxLayout()
        self.cell=QLineEdit(); self.cell.setPlaceholderText("Cell (e.g., B12)")
        self.label=QLineEdit(); self.label.setPlaceholderText("Label (e.g., Turnover)")
        add=QPushButton("Add mapping"); add.clicked.connect(self.add)
        row.addWidget(self.cell); row.addWidget(self.label); row.addWidget(add)
        layout.addLayout(row)
        self.list=QListWidget(); layout.addWidget(self.list)
        save=QPushButton("Save template JSON"); save.clicked.connect(self.save); layout.addWidget(save)
        self.setLayout(layout); self.resize(720,420)

    def load(self):
        p,_=QFileDialog.getOpenFileName(self,"Select .xlsx","","Excel (*.xlsx)")
        if not p: return
        self.path=p
        self.wb=load_workbook(p, data_only=True)
        self.sheet=self.wb.sheetnames[0]
        self.info.setText(f"Loaded: {p}\nSheet: {self.sheet} (local version captures single-cell mappings)")

    def add(self):
        if not self.wb: return
        cell=self.cell.text().strip().upper()
        label=self.label.text().strip()
        if not cell or not label: return
        self.mappings.append({"label": label, "sheet": self.sheet, "range": {"start":{"c": ''.join([x for x in cell if x.isalpha()]), "r": int(''.join([x for x in cell if x.isdigit()]))}, "end":{"c": ''.join([x for x in cell if x.isalpha()]), "r": int(''.join([x for x in cell if x.isdigit()]))}}})
        ws=self.wb[self.sheet]; val=ws[cell].value
        self.list.addItem(f"{label} -> {cell} ({val})")
        self.cell.setText(""); self.label.setText("")

    def save(self):
        name=self.template.text().strip() or "Template"
        tid=str(uuid.uuid4())
        obj={"template_id":tid,"name":name,"mappings":self.mappings,"created_local":True}
        p,_=QFileDialog.getSaveFileName(self,"Save template JSON",f"{name}.json","JSON (*.json)")
        if not p: return
        with open(p,"w",encoding="utf-8") as f: json.dump(obj,f,indent=2)
        self.info.setText(f"Saved: {p}")

if __name__=="__main__":
    app=QApplication(sys.argv)
    w=App(); w.show()
    sys.exit(app.exec())
