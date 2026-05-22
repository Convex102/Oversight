from __future__ import annotations

"""
ASGL Template Picker Service (v13.5)

Purpose
-------
Supports "Excel-like" template mapping:
- user uploads a working paper workbook
- UI shows a sheet grid preview
- user selects one or more cell ranges
- user assigns labels to each selection
- mappings are saved as a reusable Template

Storage
-------
This service stores:
- uploaded workbooks (for grid preview) under ./asgl_outputs/template_picker/<tenant_id>/uploads
- saved template JSON under ./asgl_outputs/template_picker/<tenant_id>/templates

This module is used by FastAPI endpoints in app.main.

Security
--------
FastAPI should provide tenant_id isolation (do NOT use these functions across tenants).
"""

import json
import uuid
import datetime as dt
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook


# -----------------------------
# Paths
# -----------------------------

def _root() -> Path:
    p = Path("./asgl_outputs/template_picker").resolve()
    p.mkdir(parents=True, exist_ok=True)
    return p


def _tenant_dir(tenant_id: str) -> Path:
    d = _root() / str(tenant_id)
    (d / "uploads").mkdir(parents=True, exist_ok=True)
    (d / "templates").mkdir(parents=True, exist_ok=True)
    return d


def _uploads_dir(tenant_id: str) -> Path:
    return _tenant_dir(tenant_id) / "uploads"


def _templates_dir(tenant_id: str) -> Path:
    return _tenant_dir(tenant_id) / "templates"


def _safe_name(s: str) -> str:
    out = []
    for ch in (s or ""):
        if ch.isalnum() or ch in ("-", "_", ".", " "):
            out.append(ch)
        else:
            out.append("_")
    return "".join(out).strip() or "file"


def _a1_cols(n: int) -> List[str]:
    # only render up to 26 columns for web preview simplicity
    cols = []
    for i in range(1, n + 1):
        cols.append(chr(ord("A") + i - 1))
    return cols


# -----------------------------
# Core operations
# -----------------------------

def save_upload(data: bytes, filename: str, *, tenant_id: str = "public") -> Dict[str, Any]:
    """
    Persist an uploaded workbook for preview.

    Returns:
      { upload_id, filename, saved_path, sheets }
    """
    upload_id = uuid.uuid4().hex
    name = _safe_name(filename)
    path = _uploads_dir(tenant_id) / f"{upload_id}__{name}"
    path.write_bytes(data)

    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = wb.sheetnames
    wb.close()

    return {
        "upload_id": upload_id,
        "filename": filename,
        "saved_path": str(path),
        "sheets": sheets,
        "created_at": dt.datetime.utcnow().isoformat() + "Z",
    }


def _find_upload(tenant_id: str, upload_id: str) -> Path:
    base = _uploads_dir(tenant_id)
    matches = list(base.glob(f"{upload_id}__*"))
    if not matches:
        raise FileNotFoundError("upload_id not found")
    return matches[0]


def load_sheet_grid(upload_id: str, sheet: str, *, tenant_id: str = "public",
                    max_rows: int = 30, max_cols: int = 18) -> Dict[str, Any]:
    """
    Return a lightweight grid preview for UI.
    """
    path = _find_upload(tenant_id, upload_id)
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError("sheet not found")
    ws = wb[sheet]

    rows = list(range(1, max_rows + 1))
    cols = _a1_cols(max_cols)

    values: List[List[Any]] = []
    for r in rows:
        row = []
        for ci, c in enumerate(cols, start=1):
            try:
                v = ws.cell(row=r, column=ci).value
            except Exception:
                v = None
            if isinstance(v, str):
                v = v.strip()
                if len(v) > 40:
                    v = v[:37] + "…"
            row.append(v)
        values.append(row)

    wb.close()

    return {"rows": rows, "cols": cols, "values": values}


def save_template(name: str, upload_id: str, mappings: List[Dict[str, Any]], *, tenant_id: str = "public") -> str:
    """
    Store template mapping JSON and return template_id.
    """
    template_id = uuid.uuid4().hex
    payload = {
        "template_id": template_id,
        "name": (name or "template").strip(),
        "upload_id": upload_id,
        "mappings": mappings or [],
        "created_at": dt.datetime.utcnow().isoformat() + "Z",
    }
    path = _templates_dir(tenant_id) / f"{template_id}.json"
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return template_id


def load_template(template_id: str, *, tenant_id: str = "public") -> Dict[str, Any]:
    path = _templates_dir(tenant_id) / f"{template_id}.json"
    if not path.exists():
        raise FileNotFoundError("template not found")
    return json.loads(path.read_text(encoding="utf-8"))


def list_templates(*, tenant_id: str = "public") -> List[Dict[str, Any]]:
    out = []
    for p in sorted(_templates_dir(tenant_id).glob("*.json")):
        try:
            out.append(json.loads(p.read_text(encoding="utf-8")))
        except Exception:
            continue
    # newest first
    out.sort(key=lambda d: d.get("created_at", ""), reverse=True)
    return out


# -----------------------------
# Helpers for pipeline integration
# -----------------------------

def _normalize(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s


def extract_values(upload_id: str, sheet: str, rng: dict, *, tenant_id: str = "public") -> Dict[str, Any]:
    """Return a compact preview of values in selected range."""
    grid = load_sheet_grid(
        upload_id,
        sheet,
        tenant_id=tenant_id,
        max_rows=max(int(rng["end"]["r"]), int(rng["start"]["r"]), 30),
        max_cols=26,
    )
    cols = grid["cols"]
    rows = grid["rows"]
    vals = grid["values"]
    c1 = cols.index(rng["start"]["c"])
    c2 = cols.index(rng["end"]["c"])
    r1 = rows.index(int(rng["start"]["r"]))
    r2 = rows.index(int(rng["end"]["r"]))
    rmin, rmax = min(r1, r2), max(r1, r2)
    cmin, cmax = min(c1, c2), max(c1, c2)

    preview = []
    for ri in range(rmin, min(rmax + 1, rmin + 12)):
        row = []
        for ci in range(cmin, min(cmax + 1, cmin + 10)):
            row.append(_normalize(vals[ri][ci]))
        preview.append(row)
    return {"shape": [rmax - rmin + 1, cmax - cmin + 1], "preview": preview}


KEYWORDS = [
    "turnover", "revenue", "sales", "gross profit", "profit", "operating profit", "ebit", "ebitda",
    "balance sheet", "assets", "liabilities", "equity", "cash", "bank", "vat", "vat number",
    "company number", "companies house", "trial balance", "tb", "director", "share capital",
    "notes", "depreciation", "amortisation", "dividends", "tax", "corporation tax", "ct",
    "sign-off", "review", "period", "year", "date",
]


def suggest_cells(upload_id: str, sheet: str, query: str, *, tenant_id: str = "public", limit: int = 12) -> List[Dict[str, Any]]:
    """Naive keyword scan over first N rows/cols to suggest likely labels."""
    path = _find_upload(tenant_id, upload_id)
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]

    q = (query or "").lower().strip()
    targets = [q] if q else KEYWORDS

    hits: List[Dict[str, Any]] = []
    max_rows = 80
    max_cols = 18
    for r in range(1, max_rows + 1):
        for ci in range(1, max_cols + 1):
            v = ws.cell(row=r, column=ci).value
            if not v:
                continue
            s = str(v).lower()
            for t in targets:
                if t and t in s:
                    col = chr(ord("A") + ci - 1) if ci <= 26 else f"C{ci}"
                    hits.append({"cell": f"{col}{r}", "text": str(v)[:160]})
                    break
            if len(hits) >= limit:
                wb.close()
                return hits

    wb.close()
    return hits


def apply_template(upload_id: str, template: dict, *, tenant_id: str = "public") -> Dict[str, Any]:
    """
    Extract mapped values for a saved template.

    Output format:
      { "<label>": <scalar or preview>, ... }
    """
    path = _find_upload(tenant_id, upload_id)
    wb = load_workbook(path, data_only=True, read_only=True)
    out: Dict[str, Any] = {}

    for m in (template.get("mappings") or []):
        try:
            label = str(m.get("label") or "").strip() or "field"
            sheet = m.get("sheet") or wb.sheetnames[0]
            rng = m.get("range") or {}
            ws = wb[sheet]

            # If range is missing, skip
            if not rng or "start" not in rng or "end" not in rng:
                continue

            # Determine range
            c1 = rng["start"]["c"]
            r1 = int(rng["start"]["r"])
            c2 = rng["end"]["c"]
            r2 = int(rng["end"]["r"])

            # If single-cell, return scalar
            if c1 == c2 and r1 == r2:
                v = ws[f"{c1}{r1}"].value
                out[label] = _normalize(v)
            else:
                # return a compact preview for multi-cell selections
                out[label] = extract_values(upload_id, sheet, rng, tenant_id=tenant_id)
        except Exception:
            continue

    wb.close()
    return out


def _col_to_index(col: str) -> int:
    col = (col or "").strip().upper()
    if not col:
        return 0
    n = 0
    for ch in col:
        if not ("A" <= ch <= "Z"):
            return 0
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n


def _index_to_col(idx: int) -> str:
    if idx <= 0:
        return ""
    s = ""
    n = idx
    while n > 0:
        n, r = divmod(n - 1, 26)
        s = chr(ord("A") + r) + s
    return s


def preview_mappings(
    upload_id: str,
    mappings: List[Dict[str, Any]],
    *,
    tenant_id: str = "public",
    max_cells: int = 4000,
) -> Dict[str, Any]:
    """Preview extracted values for a set of mappings against the uploaded workbook.

    This allows users to verify that the label → cell-range mapping returns the expected value(s)
    BEFORE saving a template or running a full ASGL job.

    Mapping shape (recommended):
      {
        "label": "Turnover",
        "sheet": "TB",
        "range": {"start":{"c":"B","r":12}, "end":{"c":"B","r":12}},
        "agg": "auto|cell|sum|text_join|list",
      }

    Returns:
      {
        "ok": True,
        "preview": [
          {"label": "...", "sheet":"...", "a1":"B12:B12", "agg":"cell", "value": 1234.0, "notes":""},
          ...
        ]
      }
    """
    path = _find_upload(tenant_id, upload_id)
    wb = load_workbook(path, data_only=True, read_only=True)

    out: List[Dict[str, Any]] = []
    for m in (mappings or [])[:250]:
        label = str(m.get("label") or "").strip() or "variable"
        sheet = str(m.get("sheet") or "").strip()
        agg = str(m.get("agg") or "auto").strip().lower()
        rr = m.get("range") or {}
        st = rr.get("start") or {}
        en = rr.get("end") or st

        try:
            if sheet not in wb.sheetnames:
                raise ValueError("sheet not found")

            ws = wb[sheet]
            c1 = _col_to_index(st.get("c"))
            c2 = _col_to_index(en.get("c"))
            r1 = int(st.get("r"))
            r2 = int(en.get("r"))
            if c1 <= 0 or c2 <= 0 or r1 <= 0 or r2 <= 0:
                raise ValueError("invalid range")

            # normalize ordering
            c_min, c_max = min(c1, c2), max(c1, c2)
            r_min, r_max = min(r1, r2), max(r1, r2)

            # safety bound
            cells = (c_max - c_min + 1) * (r_max - r_min + 1)
            if cells > max_cells:
                raise ValueError(f"range too large ({cells} cells)")

            values: List[Any] = []
            for r in range(r_min, r_max + 1):
                for c in range(c_min, c_max + 1):
                    try:
                        values.append(ws.cell(row=r, column=c).value)
                    except Exception:
                        values.append(None)

            # aggregation
            notes = ""
            if agg == "auto":
                agg = "cell" if len(values) == 1 else "sum"

            if agg == "cell":
                value = _normalize(values[0] if values else None)
            elif agg == "sum":
                nums = []
                for v in values:
                    try:
                        if v is None or v == "":
                            continue
                        nums.append(float(v))
                    except Exception:
                        continue
                value = float(sum(nums)) if nums else 0.0
                notes = f"summed {len(nums)} numeric cell(s)"
            elif agg == "text_join":
                parts = [str(v).strip() for v in values if v is not None and str(v).strip()]
                value = " | ".join(parts[:50])
                if len(parts) > 50:
                    notes = "truncated join"
            elif agg == "list":
                value = [_normalize(v) for v in values[:200]]
                if len(values) > 200:
                    notes = "truncated list"
            else:
                value = _normalize(values[0] if values else None)
                notes = f"unknown agg '{agg}', defaulted to cell"

            a1 = f"{_index_to_col(c_min)}{r_min}:{_index_to_col(c_max)}{r_max}"
            out.append({"label": label, "sheet": sheet, "a1": a1, "agg": agg, "value": value, "notes": notes})
        except Exception as e:
            out.append({"label": label, "sheet": sheet, "a1": "", "agg": agg, "error": str(e), "value": None, "notes": ""})

    wb.close()
    return {"ok": True, "preview": out}

# -----------------------------
# Direct workbook path helpers (worker-friendly)
# -----------------------------

def apply_mappings_to_path(workbook_path: str, mappings: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Apply a mapping list directly to a workbook path.

    This is the worker-friendly equivalent of apply_template(upload_id, template).
    It returns a dict keyed by mapping label, including evidence metadata.

    Output example:
      {
        "Turnover": {"value": 12345.0, "sheet":"TB", "range":"B12", "aggregation":"single"},
        "Directors loan": {"value": {"shape":[..],"preview":[..]}, "sheet":"WP", "range":"C5:D10", "aggregation":"preview"}
      }
    """
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    path = Path(workbook_path)
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    wb = load_workbook(str(path), data_only=True, read_only=True)
    out: Dict[str, Any] = {}

    def _cell_addr(c: str, r: int) -> str:
        return f"{c}{int(r)}"

    for m in (mappings or []):
        try:
            label = str(m.get("label") or "").strip() or "field"
            sheet = m.get("sheet") or wb.sheetnames[0]
            rng = m.get("range") or {}
            agg = (m.get("aggregation") or "single").strip().lower()

            if sheet not in wb.sheetnames:
                sheet = wb.sheetnames[0]
            ws = wb[sheet]

            if not rng or "start" not in rng or "end" not in rng:
                continue

            c1 = str(rng["start"]["c"])
            r1 = int(rng["start"]["r"])
            c2 = str(rng["end"]["c"])
            r2 = int(rng["end"]["r"])
            rmin, rmax = min(r1, r2), max(r1, r2)
            cmin, cmax = min(column_index_from_string(c1), column_index_from_string(c2)), max(column_index_from_string(c1), column_index_from_string(c2))

            # single cell
            if rmin == rmax and cmin == cmax:
                addr = _cell_addr(c1, r1)
                v = ws[addr].value
                out[label] = {
                    "value": _normalize(v),
                    "sheet": sheet,
                    "range": addr,
                    "aggregation": "single",
                }
                continue

            # multi-range: compute aggregates or preview
            values = []
            # hard cap to avoid huge ranges in worker
            rcap = min(rmax, rmin + 250)
            ccap = min(cmax, cmin + 50)
            for rr in range(rmin, rcap + 1):
                row_vals = []
                for cc in range(cmin, ccap + 1):
                    v = ws.cell(row=rr, column=cc).value
                    row_vals.append(_normalize(v))
                values.append(row_vals)

            flat = [v for row in values for v in row if isinstance(v, (int, float))]

            if agg in ("sum", "total") and flat:
                out[label] = {"value": float(sum(flat)), "sheet": sheet, "range": f"{c1}{r1}:{c2}{r2}", "aggregation": "sum"}
            elif agg in ("avg", "mean") and flat:
                out[label] = {"value": float(sum(flat) / max(1, len(flat))), "sheet": sheet, "range": f"{c1}{r1}:{c2}{r2}", "aggregation": "avg"}
            elif agg in ("min",) and flat:
                out[label] = {"value": float(min(flat)), "sheet": sheet, "range": f"{c1}{r1}:{c2}{r2}", "aggregation": "min"}
            elif agg in ("max",) and flat:
                out[label] = {"value": float(max(flat)), "sheet": sheet, "range": f"{c1}{r1}:{c2}{r2}", "aggregation": "max"}
            elif agg in ("first",):
                out[label] = {"value": values[0][0] if values and values[0] else None, "sheet": sheet, "range": f"{c1}{r1}:{c2}{r2}", "aggregation": "first"}
            else:
                # default: compact preview (matches extract_values output)
                out[label] = {
                    "value": {"shape": [rmax - rmin + 1, cmax - cmin + 1], "preview": [row[:10] for row in values[:12]]},
                    "sheet": sheet,
                    "range": f"{c1}{r1}:{c2}{r2}",
                    "aggregation": "preview",
                }
        except Exception:
            continue

    try:
        wb.close()
    except Exception:
        pass
    return out
